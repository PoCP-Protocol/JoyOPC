from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


HEADER_ALIASES = {
    "product_name": ["product_name", "name", "product", "商品名称", "产品名称", "品名", "标题"],
    "supplier_name": ["supplier_name", "supplier", "供应商", "供应商名称", "厂家", "工厂"],
    "supplier_sku": ["supplier_sku", "sku", "货号", "供应商sku", "型号", "model"],
    "supplier_price": ["supplier_price", "price", "cost", "供货价", "供应价", "出厂价", "采购价"],
    "estimated_landed_cost": ["estimated_landed_cost", "landed_cost", "落地成本", "到岸成本", "综合成本"],
    "target_retail_price": ["target_retail_price", "retail_price", "sale_price", "零售价", "建议零售价", "售价"],
    "market": ["market", "市场", "国家", "country"],
    "recommended_channel": ["recommended_channel", "channel", "渠道", "平台", "销售平台"],
    "exclusive_rights": ["exclusive_rights", "exclusive", "独家", "独家权", "排他销售权"],
    "uniqueness": ["uniqueness", "独特性", "差异化"],
    "channel_control": ["channel_control", "渠道控制", "渠道控制力"],
    "cost_advantage": ["cost_advantage", "成本优势"],
    "supply_advantage": ["supply_advantage", "供应优势"],
    "content_advantage": ["content_advantage", "内容优势"],
    "brand_advantage": ["brand_advantage", "品牌优势"],
    "compliance_risk": ["compliance_risk", "合规风险"],
    "return_risk": ["return_risk", "退货风险"],
    "cash_cycle_days": ["cash_cycle_days", "周转天数", "资金周转天数"],
    "has_persona": ["has_persona", "人格", "有人格"],
    "memory_enabled": ["memory_enabled", "记忆", "记忆开启"],
    "soul_recipe_id": ["soul_recipe_id", "魂配方", "soul"],
    "hw_gen": ["hw_gen", "硬件代际", "gen"],
    "module_tier": ["module_tier", "模组档位"],
    "shell": ["shell", "外形壳", "外形"],
    "claimed_features": ["claimed_features", "硬件特征", "features"],
    "certs_held": ["certs_held", "已有认证", "认证"],
    "skills": ["skills", "技能"],
}

MARKET_ALIASES = {
    "source": ["source", "数据源", "来源"],
    "market": ["market", "市场", "国家", "country"],
    "channel": ["channel", "渠道", "平台"],
    "keyword": ["keyword", "query", "关键词", "搜索词", "品类词"],
    "demand_score": ["demand_score", "demand", "需求分", "需求指数", "搜索热度"],
    "growth_score": ["growth_score", "growth", "增长分", "增长指数", "趋势"],
    "social_velocity": ["social_velocity", "velocity", "社交传播", "内容热度"],
    "competition_score": ["competition_score", "competition", "竞争分", "竞争强度"],
    "median_price": ["median_price", "price", "中位价", "中位价格", "平均售价"],
    "confidence": ["confidence", "置信度", "confidence_score"],
}


def _norm(s: Any) -> str:
    return re.sub(r"[\s_\-（）()]+", "", str(s or "").strip().lower())


def _header_map(headers: Iterable[Any], aliases: dict[str, list[str]]) -> dict[str, str]:
    normalized = {_norm(h): str(h) for h in headers if h is not None}
    result: dict[str, str] = {}
    for canonical, choices in aliases.items():
        for choice in choices:
            if _norm(choice) in normalized:
                result[canonical] = normalized[_norm(choice)]
                break
    return result


def _float(v: Any, default: float = 0.0) -> float:
    if v is None or str(v).strip() == "":
        return default
    s = str(v).strip().replace(",", "")
    s = re.sub(r"^[¥$€£]", "", s)
    if s.endswith("%"):
        s = s[:-1]
    try:
        return float(s)
    except ValueError:
        return default


def _int(v: Any, default: int = 0) -> int:
    return int(round(_float(v, float(default))))


def _bool(v: Any) -> bool:
    return str(v or "").strip().lower() in {"1", "true", "yes", "y", "是", "有", "独家", "exclusive"}


def _read_tabular(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "utf-8", "gb18030"):
            try:
                text = raw.decode(encoding)
                return list(csv.DictReader(io.StringIO(text)))
            except UnicodeDecodeError:
                continue
        raise ValueError("CSV encoding is not supported")
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return []
        return [dict(zip(headers, row)) for row in rows if any(v is not None and str(v).strip() for v in row)]
    raise ValueError("Only CSV/XLSX/XLSM supplier files are supported")


@dataclass(slots=True)
class SupplierImportRecord:
    normalized: dict[str, Any]
    raw: dict[str, Any]


class SupplierCatalogImporter:
    def parse(self, path: Path) -> list[SupplierImportRecord]:
        rows = _read_tabular(path)
        if not rows:
            return []
        mapping = _header_map(rows[0].keys(), HEADER_ALIASES)
        if "product_name" not in mapping:
            raise ValueError("Supplier file must contain a product name column")

        out: list[SupplierImportRecord] = []
        for row in rows:
            get = lambda key, default=None: row.get(mapping[key], default) if key in mapping else default
            product_name = str(get("product_name", "")).strip()
            if not product_name:
                continue
            supplier_price = _float(get("supplier_price"))
            landed = _float(get("estimated_landed_cost"), supplier_price * 1.35 if supplier_price else 0)
            retail = _float(get("target_retail_price"), landed * 2.4 if landed else 0)
            normalized = {
                "product_name": product_name,
                "source": "Supplier File Import",
                "supplier_name": str(get("supplier_name", "Imported Supplier") or "Imported Supplier").strip(),
                "supplier_sku": str(get("supplier_sku", "") or "").strip(),
                "market": str(get("market", "US") or "US").strip().upper(),
                "recommended_channel": str(get("recommended_channel", "TikTok Shop") or "TikTok Shop").strip(),
                "supplier_price": supplier_price,
                "estimated_landed_cost": landed,
                "target_retail_price": retail,
                "exclusive_rights": _bool(get("exclusive_rights")),
                "uniqueness": _float(get("uniqueness"), 50),
                "channel_control": _float(get("channel_control"), 50),
                "cost_advantage": _float(get("cost_advantage"), 50),
                "supply_advantage": _float(get("supply_advantage"), 55),
                "content_advantage": _float(get("content_advantage"), 50),
                "brand_advantage": _float(get("brand_advantage"), 40),
                "compliance_risk": _float(get("compliance_risk"), 35),
                "return_risk": _float(get("return_risk"), 35),
                "cash_cycle_days": _int(get("cash_cycle_days"), 30),
                "soul_recipe_id": str(get("soul_recipe_id", "") or "").strip(),
                "has_persona": _bool(get("has_persona")),
                "memory_enabled": _bool(get("memory_enabled")),
                "skills": str(get("skills", "") or ""),
                "hw_gen": max(1, min(4, _int(get("hw_gen"), 1) or 1)),
                "module_tier": str(get("module_tier", "Mini") or "Mini").strip(),
                "shell": str(get("shell", "") or "").strip(),
                "claimed_features": str(get("claimed_features", "") or ""),
                "certs_held": str(get("certs_held", "") or ""),
            }
            out.append(SupplierImportRecord(normalized=normalized, raw={str(k): v for k, v in row.items()}))
        return out


class MarketSignalFileImporter:
    def parse(self, path: Path) -> list[dict[str, Any]]:
        rows = _read_tabular(path)
        if not rows:
            return []
        mapping = _header_map(rows[0].keys(), MARKET_ALIASES)
        if "keyword" not in mapping:
            raise ValueError("Market file must contain keyword/search term column")
        out = []
        for row in rows:
            get = lambda key, default=None: row.get(mapping[key], default) if key in mapping else default
            keyword = str(get("keyword", "")).strip()
            if not keyword:
                continue
            out.append({
                "source": str(get("source", "Imported Market Export") or "Imported Market Export").strip(),
                "market": str(get("market", "US") or "US").strip().upper(),
                "channel": str(get("channel", "MULTI") or "MULTI").strip(),
                "keyword": keyword,
                "demand_score": max(0, min(100, _float(get("demand_score"), 50))),
                "growth_score": max(0, min(100, _float(get("growth_score"), 50))),
                "social_velocity": max(0, min(100, _float(get("social_velocity"), 50))),
                "competition_score": max(0, min(100, _float(get("competition_score"), 50))),
                "median_price": max(0, _float(get("median_price"), 0)),
                "confidence": max(0, min(100, _float(get("confidence"), 75))),
            })
        return out


def json_dumps(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)
